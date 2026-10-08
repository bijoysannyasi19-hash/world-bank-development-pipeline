import os
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def run_report():
    doc = Document()
    res = "results"
    fig = "figures"
    rep = "report"
    doc.add_heading("Integrative Capstone: Data Science Pipeline", 0)
    p = doc.add_paragraph("[Your Name]\n")
    p.add_run("Date: 2026-10-09")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    
    doc.add_heading("Table of Contents", 1)
    t = [
        "1. Title Page", "2. Executive Summary", "3. Problem Statement",
        "4. Methodological Approach", "5. Data Collection", "6. Data Cleaning and Preprocessing",
        "7. Exploratory Data Analysis", "8. Modeling", "8.1 Unsupervised Clustering",
        "8.2 Supervised Modeling", "9. Evaluation", "10. Insights and Recommendations",
        "11. Reflection on Challenges and Successes", "12. Conclusion and Further Analysis",
        "13. References", "Appendix A: Decision log", "Appendix B: Architecture diagram",
        "Appendix C: Requirement coverage table"
    ]
    for i in t:
        doc.add_paragraph(i)
    doc.add_page_break()
    
    with open(os.path.join(res, "test_results.json"), "r") as f:
        tr = json.load(f)
    with open(os.path.join(res, "cluster_metrics.json"), "r") as f:
        cm = json.load(f)
    
    rm = tr["RMSE"]
    ci = tr["RMSE_CI_95"]
    dm = tr["Dummy_RMSE"]
    k_opt = cm["k_opt"]
    
    doc.add_heading("2. Executive Summary", 1)
    doc.add_paragraph(f"I built a full data science pipeline to predict life expectancy and cluster countries. "
                      f"Using World Bank data, I identified {k_opt} distinct country profiles. "
                      f"The predictive model achieved an RMSE of {rm:.2f} years (95% CI: {ci[0]:.2f}-{ci[1]:.2f}), "
                      f"improving significantly on the baseline RMSE of {dm:.2f}. "
                      f"Healthcare spending and electricity access were key drivers.")
    
    doc.add_heading("3. Problem Statement", 1)
    doc.add_paragraph("The goal is to segment countries into development profiles and predict life expectancy. "
                      "Questions: 1. What are the natural country profiles? 2. Which economic indicators best predict life expectancy? 3. Where does the model fail? "
                      "Policy analysts will use this to target interventions. A useful answer includes actionable segments and reliable predictions.")
    
    doc.add_heading("4. Methodological Approach", 1)
    doc.add_paragraph("I acquired data via API, cleaned it by averaging over 3 years, explored it visually, "
                      "engineered features, clustered using K-Means, and trained a Gradient Boosting model. "
                      "I ordered steps this way to ensure clean data before any modeling.")
    
    doc.add_heading("5. Data Collection", 1)
    doc.add_paragraph("I fetched data from the World Bank API (CC BY 4.0). Indicators include GDP, health spending, and electricity access.")
    doc.add_paragraph("Code example for data fetching:")
    p = doc.add_paragraph("res = requests.get(url)\ndata = res.json()")
    p.runs[0].font.name = "Courier New"
    doc.add_paragraph("This code fetches JSON data from the API endpoint.")
    
    doc.add_heading("6. Data Cleaning and Preprocessing", 1)
    doc.add_paragraph("I dropped aggregates and handled missing values.")
    doc.add_paragraph("Code example:")
    p = doc.add_paragraph("dfm = dfm[na_c <= 0.3]")
    p.runs[0].font.name = "Courier New"
    doc.add_paragraph("I dropped rows with too many missing values to preserve quality.")
    
    doc.add_heading("7. Exploratory Data Analysis", 1)
    doc.add_paragraph("I analyzed distributions and correlations.")
    f_l = ["02_missingness.png", "03_distributions.png", "04_correlation_heatmap.png", "05_target_relationships.png", "06_group_comparisons.png"]
    for i, fi in enumerate(f_l):
        fp = os.path.join(fig, fi)
        if os.path.exists(fp):
            doc.add_picture(fp, width=Inches(5))
            doc.add_paragraph(f"Figure {i+1}: {fi}")
    
    doc.add_heading("8. Modeling", 1)
    doc.add_heading("8.1 Unsupervised Clustering", 2)
    doc.add_paragraph(f"I used K-Means with k={k_opt} based on the elbow method and silhouette score.")
    f_l2 = ["07_k_selection.png", "08_dendrogram.png", "09_pca_clusters.png", "10_cluster_profiles.png"]
    for i, fi in enumerate(f_l2):
        fp = os.path.join(fig, fi)
        if os.path.exists(fp):
            doc.add_picture(fp, width=Inches(5))
            doc.add_paragraph(f"Figure {i+6}: {fi}")
            
    doc.add_heading("8.2 Supervised Modeling", 2)
    doc.add_paragraph("I trained Ridge, Random Forest, and Gradient Boosting. Gradient Boosting performed best.")
    f_l3 = ["11_model_comparison.png", "12_tuning_results.png"]
    for i, fi in enumerate(f_l3):
        fp = os.path.join(fig, fi)
        if os.path.exists(fp):
            doc.add_picture(fp, width=Inches(5))
            doc.add_paragraph(f"Figure {i+10}: {fi}")
            
    doc.add_heading("9. Evaluation", 1)
    doc.add_paragraph(f"The test RMSE was {rm:.2f}, baseline {dm:.2f}.")
    f_l4 = ["13_predicted_vs_actual.png", "14_residuals.png", "15_feature_importance.png", "16_error_by_cluster.png"]
    for i, fi in enumerate(f_l4):
        fp = os.path.join(fig, fi)
        if os.path.exists(fp):
            doc.add_picture(fp, width=Inches(5))
            doc.add_paragraph(f"Figure {i+12}: {fi}")
            
    doc.add_heading("10. Insights and Recommendations", 1)
    doc.add_paragraph("1. Healthcare spending drives life expectancy. 2. Infrastructure (electricity) is crucial. "
                      "3. Some clusters need basic sanitation. 4. High GDP alone does not guarantee long life.")
    doc.add_paragraph("Recommendations: NGOs should focus on basic infrastructure in Cluster 0 countries.")
    
    doc.add_heading("11. Reflection on Challenges and Successes", 1)
    doc.add_paragraph("Data merging was tricky due to missing country codes. Dropping sparse columns helped. "
                      "The pipeline successfully grouped countries logically.")
    
    doc.add_heading("12. Conclusion and Further Analysis", 1)
    doc.add_paragraph("The model works well but fails on extreme outliers. Future work should use time-series causal modeling.")
    
    doc.add_heading("13. References", 1)
    doc.add_paragraph("World Bank Open Data: https://data.worldbank.org/")
    
    doc.add_page_break()
    doc.add_heading("Appendix A: Decision log", 1)
    dl = pd.read_csv(os.path.join(res, "decision_log.csv"))
    for idx, r in dl.iterrows():
        doc.add_paragraph(f"Decision: {r['decision']}, Choice: {r['choice']}, Reason: {r['reason']}")
        
    doc.add_heading("Appendix B: Architecture diagram", 1)
    fp = os.path.join("docs", "architecture.png")
    if os.path.exists(fp):
        doc.add_picture(fp, width=Inches(6))
        
    doc.add_heading("Appendix C: Requirement coverage table", 1)
    doc.add_paragraph("Requirement 1: Section 4-9\nRequirement 2: Full report\nRequirement 3: This docx")
    
    try:
        doc.save(os.path.join(rep, "Capstone_Report.docx"))
    except PermissionError:
        doc.save(os.path.join(rep, "Capstone_Report_Optimized.docx"))
        print("Saved as Capstone_Report_Optimized.docx due to file lock.")
    print("Report generated")

if __name__ == "__main__":
    run_report()
